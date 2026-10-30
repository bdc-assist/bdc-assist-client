# /// script
# requires-python = ">=3.12"
# dependencies = ["python-pptx"]
# ///
"""The progress deck: slides/ui-progress.pptx.

    uv run slides/build.py            # the deck
    uv run slides/build.py --preview  # also rough PNGs of each slide in slides/preview/ (for checking layout)

Every shape goes through the drawing helpers below, which add it to the deck and to the preview. The preview is
drawn with Liberation Sans, Arial's metrics, so text that fits there fits in PowerPoint. Diagrams are native
shapes, so they stay editable. Screenshots go in slides/shots/ under the name a slide's placeholder shows; the
capture instructions are in each slide's speaker notes.
"""

import sys
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
OUT = HERE / "ui-progress.pptx"
SHOTS = HERE / "shots"
PREVIEW = HERE / "preview"

FONT = "Arial"
W, H = 13.333, 7.5

# RENCI colours (design/favicon.html)
TWILIGHT = "14385F"
BLUE = "00758D"  # RENCI Blue
ULTRA = "00A9C4"  # Ultramarine
VIOLET = "6E5FA6"  # not RENCI's: a muted violet for the fourth repo (Quantum Pink stood out too much)
FROST = "D6E8F2"
STONE = "666676"
SLATE = "F3F3F8"
WHITE = "FFFFFF"
INK = "22263A"  # body text

# one accent per repo, the same on every diagram; services run elsewhere are dashed stone
REPO = {"client": TWILIGHT, "assist": BLUE, "mcp": ULTRA, "builder": VIOLET, "external": STONE}


def tint(hex_, amount):
    """`hex_` mixed with white: amount 0 = white, 1 = the colour."""
    c = [int(hex_[i : i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(255 - (255 - v) * amount):02X}" for v in c)


def rgb(hex_):
    return RGBColor.from_string(hex_)


# ---- drawing: each helper adds to the slide and records the shape for the preview -------------------------


class Slide:
    def __init__(self, deck, bg=WHITE):
        self.s = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])  # blank
        self.ops = [("bg", bg)]
        self.bg = bg
        fill = self.s.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(bg)
        deck.slides.append(self)

    def notes(self, text):
        self.s.notes_slide.notes_text_frame.text = text.strip()

    def text(self, x, y, w, h, runs, size=14, color=INK, bold=False, align="l", anchor="t", italic=False, spacing=1.0):
        """`runs`: a string (\\n starts a paragraph) or a list of paragraphs, each a list of (text, options)."""
        tb = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        self._fill_text(tb.text_frame, runs, size, color, bold, align, anchor, italic, spacing, margin=0)
        self.ops.append(("text", x, y, w, h, _paras(runs, size, color, bold, italic), align, anchor, 0, spacing))
        return tb

    def box(self, x, y, w, h, fill=WHITE, line=None, width=1.5, dashed=False, radius=0.08, shape="round"):
        kind = {"round": MSO_SHAPE.ROUNDED_RECTANGLE, "rect": MSO_SHAPE.RECTANGLE, "oval": MSO_SHAPE.OVAL}[shape]
        sh = self.s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        if shape == "round":
            sh.adjustments[0] = min(0.5, radius / min(w, h))
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = rgb(line)
            sh.line.width = Pt(width)
            if dashed:
                sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        self.ops.append(("box", x, y, w, h, fill, line, width, dashed, radius, shape))
        return sh

    def label_box(self, x, y, w, h, runs, size=12, color=INK, align="l", anchor="t", pad=0.12, **box):
        """A box with text inside it, padded."""
        sh = self.box(x, y, w, h, **box)
        self._fill_text(sh.text_frame, runs, size, color, False, align, anchor, False, 1.0, margin=pad)
        self.ops.append(("text", x, y, w, h, _paras(runs, size, color, False, False), align, anchor, pad, 1.0))
        return sh

    def line(self, x1, y1, x2, y2, color=STONE, width=1.5, arrow=True, dashed=False):
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(width)
        if dashed:
            c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if arrow:
            ln = c.line._get_or_add_ln()
            ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
        self.ops.append(("line", x1, y1, x2, y2, color, width, arrow, dashed))

    def path(self, points, **kw):
        """An arrow through `points`, the head on the last segment."""
        for i, (a, b) in enumerate(zip(points, points[1:])):
            self.line(*a, *b, **{**kw, "arrow": kw.get("arrow", True) and i == len(points) - 2})

    def picture(self, file, x, y, w, h):
        """`file` fitted into the box, centred."""
        from PIL import Image

        iw, ih = Image.open(file).size
        scale = min(w / iw, h / ih)
        pw, ph = iw * scale, ih * scale
        self.s.shapes.add_picture(str(file), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
        self.ops.append(("image", file, x + (w - pw) / 2, y + (h - ph) / 2, pw, ph))

    @staticmethod
    def _fill_text(tf, runs, size, color, bold, align, anchor, italic, spacing, margin):
        tf.word_wrap = True
        tf.auto_size = None
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        for i, para in enumerate(_paras(runs, size, color, bold, italic)):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            p.line_spacing = spacing
            if para and para[0][1].get("space_before"):
                p.space_before = Pt(para[0][1]["space_before"])
            for text, o in para:
                r = p.add_run()
                r.text = text
                f = r.font
                f.name = FONT
                f.size = Pt(o["size"])
                f.bold = o["bold"]
                f.italic = o["italic"]
                f.color.rgb = rgb(o["color"])


def _paras(runs, size, color, bold, italic):
    """Paragraphs of (text, options) with every option filled in."""
    base = {"size": size, "color": color, "bold": bold, "italic": italic}
    if isinstance(runs, str):
        runs = [[(line, {})] for line in runs.split("\n")]
    return [[(t, {**base, **o}) for t, o in para] for para in runs]


def b(text, **o):
    return (text, {"bold": True, **o})


def r(text, **o):
    return (text, o)


# ---- the deck's frames --------------------------------------------------------------------------------------


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.slides = []
        self.when = date.today().strftime("%B %Y")

    def content(self, title, lede=None):
        """A white slide with its title at the top, the same place on every one."""
        s = Slide(self)
        s.text(0.6, 0.4, W - 1.2, 0.7, title, size=30, color=TWILIGHT, bold=True, anchor="m")
        if lede:
            s.text(0.6, 1.1, W - 1.2, 0.45, lede, size=16, color=STONE)
        n = len(self.slides)
        s.text(0.6, 7.0, 6, 0.3, f"BDC Assist · progress · {self.when}", size=10, color=STONE)
        s.text(W - 1.6, 7.0, 1.0, 0.3, str(n), size=10, color=STONE, align="r")
        return s

    def divider(self, number, title, lede):
        s = Slide(self, bg=TWILIGHT)
        s.text(0.9, 2.3, 3, 0.6, f"{number:02d}", size=28, color=ULTRA, bold=True)
        s.text(0.9, 2.9, W - 1.8, 1.1, title, size=44, color=WHITE, bold=True)
        s.text(0.9, 4.05, W - 2.8, 1.2, lede, size=20, color=FROST)
        return s

    def save(self):
        self.prs.save(OUT)


# ---- diagram pieces -----------------------------------------------------------------------------------------


def node(s, x, y, w, h, title, body=None, repo="assist", size=13, body_size=11, align="l", anchor="t"):
    """A component: tinted fill and border in its repo's colour; dashed stone for things run elsewhere."""
    accent = REPO[repo]
    ext = repo == "external"
    runs = [[b(title, size=size, color=TWILIGHT)]]
    if body:
        runs += [[r(line, size=body_size, color=INK, space_before=3)] for line in body.split("\n")]
    return s.label_box(
        x, y, w, h, runs, align=align, anchor=anchor,
        fill=SLATE if ext else tint(accent, 0.12), line=accent, width=1.75, dashed=ext,
    )


def tag(s, x, y, w, text, align="c", size=10, color=STONE, h=0.3):
    """A small label beside an arrow."""
    s.text(x, y, w, h, text, size=size, color=color, align=align, anchor="m")


def key(s, x, y, items):
    """Colour key: (repo, name) pairs in a row."""
    for repo, name in items:
        ext = repo == "external"
        s.box(x, y + 0.04, 0.3, 0.22, fill=SLATE if ext else tint(REPO[repo], 0.12), line=REPO[repo],
              width=1.5, dashed=ext, radius=0.04)
        s.text(x + 0.4, y, 1.9, 0.3, name, size=10, color=STONE, anchor="m")
        x += 0.45 + 0.085 * len(name) + 0.25


# ---- slides -------------------------------------------------------------------------------------------------


def title_slide(d):
    s = Slide(d, bg=TWILIGHT)
    # the motif: the diagrams' rounded components, as a quiet cluster on the right
    for x, y, w, h, c in [(9.2, 1.2, 2.6, 1.3, BLUE), (10.4, 2.9, 2.3, 1.3, ULTRA), (8.6, 4.6, 2.4, 1.3, VIOLET),
                          (11.2, 4.9, 1.6, 1.0, FROST)]:
        s.box(x, y, w, h, fill=None, line=c, width=2.25)
    s.text(0.9, 2.2, 7.5, 1.2, "BDC Assist", size=54, color=WHITE, bold=True)
    s.text(0.9, 3.35, 7.5, 1.2, "Progress on the BioData Catalyst assistant and its web client", size=24, color=FROST)
    s.text(0.9, 5.4, 7.5, 0.4, d.when, size=16, color=ULTRA)
    s.notes("""
Title slide. No screenshot.
Talking points: where the assistant stands; the system first (three services and the web client), then a walk
through the web client's features, then what's next.
""")


def architecture(d):
    s = d.content("How the pieces fit together")
    # the serving side: what answers a question; the builder stays outside and only pushes finished chunks
    s.box(3.25, 1.75, 6.75, 3.25, fill=tint(FROST, 0.45), line=TWILIGHT, width=1, dashed=True, radius=0.12)
    s.text(3.45, 1.82, 3, 0.3, "Serving side", size=11, color=TWILIGHT, bold=True)

    node(s, 0.6, 2.75, 2.05, 1.5, "bdc-assist-client", "Browser app: chat, saved\nconversations, graphs", "client")
    node(s, 3.55, 2.2, 2.9, 2.45, "bdc-assist",
         "Chat API (FastAPI)\nLangGraph workflow:\nguardrails, canned answers,\nfollow-ups\nAgent that calls tools", "assist")
    node(s, 7.0, 2.95, 2.7, 1.7, "bdc-doc-mcp", "search_docs tool\nVector DB (Chroma)\nIngest API", "mcp")
    node(s, 10.4, 1.95, 2.35, 1.2, "Dug MCP", "RENCI, another team:\nconcepts, studies, variables", "external")
    node(s, 3.55, 5.4, 2.9, 1.25, "Completion LLM", "OpenAI-compatible endpoint:\nanswers and every check",
         "external")
    node(s, 7.0, 5.4, 2.7, 1.25, "bdc-doc-builder", "Reads GitBook, website,\nvideo transcripts, FAQs", "builder")
    node(s, 10.4, 5.4, 2.35, 1.25, "Embedding model", "bge-m3 on Ollama,\nthe same on both sides",
         "external")

    s.line(2.65, 3.5, 3.55, 3.5)
    tag(s, 2.6, 2.95, 1.0, "/chat/stream\n(SSE)", h=0.5)
    s.line(6.45, 3.8, 7.0, 3.8)
    tag(s, 6.4, 3.45, 0.65, "MCP")
    s.line(6.45, 2.5, 10.4, 2.5)
    tag(s, 6.6, 2.15, 3.3, "MCP, results turned into knowledge graphs")
    s.line(5.0, 4.65, 5.0, 5.4)
    s.line(8.35, 5.4, 8.35, 4.65)
    tag(s, 8.45, 5.02, 1.4, "ingest (token)", align="l")
    s.path([(9.7, 3.8), (11.57, 3.8), (11.57, 5.4)])
    tag(s, 10.05, 3.45, 1.5, "query vectors")
    s.line(9.7, 6.02, 10.4, 6.02)
    s.line(7.0, 6.02, 6.45, 6.02)

    key(s, 0.6, 6.45, [("external", "run elsewhere")])
    s.notes("""
No screenshot.
- The client talks to one endpoint, bdc-assist's /chat/stream, which streams progress, the answer and its graphs.
- bdc-assist runs a LangGraph workflow (next slide) around an agent. The agent's tools come from two MCP servers:
  bdc-doc-mcp (our doc search) and Dug's (RENCI, run by another team). An interceptor in bdc-assist turns Dug's
  rows into the knowledge graphs the client draws; the LLM never sees the graphs.
- The doc builder is deliberately outside the serving side: it scrapes, chunks and embeds, then pushes finished
  records through a token-guarded ingest API. Nothing on the serving side scrapes or writes on its own.
- One completion model does the answering and every guardrail check; one embedding model must be shared by the
  builder (documents) and doc-mcp (queries), or the vectors don't match.
- The three Python repos are the generic "r-" template (r-assist, r-doc-mcp, r-doc-builder) configured for BDC.
""")


def workflow(d):
    s = d.content("bdc-assist: the workflow around every answer")
    y, h = 2.95, 1.05  # the main row
    mid = y + h / 2
    cols = {}  # name -> (x, w)
    x = 0.6
    for name, w in [("ctx", 1.45), ("par", 1.55), ("gate", 0.85), ("agent", 1.1), ("out", 1.6), ("disc", 1.7),
                    ("fup", 1.7)]:
        cols[name] = (x, w)
        x += w + 0.35

    def at(name, title, body, yy=y, hh=h, repo="assist"):
        cx, cw = cols[name]
        node(s, cx, yy, cw, hh, title, body, repo, size=11, body_size=10)

    at("ctx", "contextualize", "follow-ups made standalone")
    at("par", "classify", "scripted topics", yy=2.0, hh=0.85)
    at("par", "input_guardrail", "policy check", yy=4.05, hh=0.85)
    at("gate", "gate", "joins the two")
    at("agent", "agent", "calls tools on the MCP servers")
    at("out", "output_guardrail", "checks the answer against the question")
    at("disc", "append_disclaimer", "flagged topics", yy=1.75, hh=0.85)
    at("fup", "suggest_followups", "up to 3 follow-up questions")

    def right(name):
        return cols[name][0] + cols[name][1]

    def left(name):
        return cols[name][0]

    def centre(name):
        return cols[name][0] + cols[name][1] / 2

    px = (right("ctx") + left("par")) / 2
    s.path([(right("ctx"), mid), (px, mid), (px, 2.42), (left("par"), 2.42)])
    s.path([(px, mid), (px, 4.47), (left("par"), 4.47)])
    gx = (right("par") + left("gate")) / 2
    s.path([(right("par"), 2.42), (gx, 2.42), (gx, mid), (left("gate"), mid)])
    s.path([(right("par"), 4.47), (gx, 4.47), (gx, mid)], arrow=False)
    s.line(right("gate"), mid, left("agent"), mid)
    s.line(right("agent"), mid, left("out"), mid)
    s.line(right("out"), mid, left("fup"), mid)
    tag(s, right("out") + 0.05, mid - 0.32, left("fup") - right("out") - 0.1, "done")
    s.path([(centre("out"), y), (centre("out"), 2.17), (left("disc"), 2.17)])
    s.path([(right("disc"), 2.17), (centre("fup"), 2.17), (centre("fup"), y)])

    # the early exits
    ex_y = 5.55
    node(s, left("gate") - 0.35, ex_y, 1.6, 0.85, "Refusal or canned answer", repo="client", size=11)
    s.line(centre("gate"), y + h, centre("gate"), ex_y, dashed=True)
    tag(s, centre("gate") + 0.1, 4.6, 1.4, "refused, canned", align="l")
    node(s, left("out"), ex_y, cols["out"][1], 0.85, "Reject reply", "blocked: output", repo="client", size=11)
    s.line(centre("out"), y + h, centre("out"), ex_y, dashed=True)
    tag(s, centre("out") + 0.1, 4.6, 1.0, "blocked", align="l")

    s.text(left("ctx"), 6.55, W - 1.2, 0.35,
           "classify and input_guardrail run side by side: one LLM round trip instead of two, about 1.6 s off every "
           "chat.", size=12, color=STONE)
    s.notes("""
No screenshot. Redrawn from bdc-assist/docs/workflow.png.
- contextualize rewrites a follow-up ("what about COPD?") into a standalone question using the chat history; no
  LLM call when there's no history.
- classify and input_guardrail run in parallel. The guardrail is a yes/no policy check (on topic, no jailbreaks or
  abuse); classify sends scripted topics to a canned answer, or flags them for a disclaimer. gate joins them; a
  blocked question gets the refusal whatever classify found.
- The agent calls tools on the MCP servers: doc search and Dug.
- output_guardrail checks the answer against the question. A "no" replaces the answer with a reject reply
  (blocked: output); the client shows that the draft was replaced.
- suggest_followups decides whether follow-ups would help and suggests up to 3 (skipped for refusals and canned
  answers).
""")


def template(d):
    s = d.content("bdc-assist: a reusable assistant, configured for BDC",
                  "The code is a generic template (r-assist); everything project-specific lives in four yaml files.")
    files = [
        ("project.yaml", "Names: project, assistant\nHow many follow-ups to suggest"),
        ("prompts.yaml", "Every prompt: the policy checklist,\nthe agent's house rules"),
        ("predefined_responses.yaml", "Canned answers and disclaimers,\nby topic, each with a link"),
        ("mcp_servers.yaml", "The MCP servers the agent calls,\nand how to use each one's tools"),
    ]
    for i, (name, body) in enumerate(files):
        x, y = 0.6 + (i % 2) * 3.75, 1.95 + (i // 2) * 1.85
        node(s, x, y, 3.5, 1.55, name, body, "assist", size=15, body_size=13)

    s.box(8.35, 1.95, 4.4, 3.4, fill=SLATE, radius=0.1)
    s.text(8.6, 2.1, 4.0, 0.4, "Configured for BDC", size=15, color=TWILIGHT, bold=True)
    items = [
        ("Interceptor: ", "turns Dug's rows into knowledge graphs and study sources"),
        ("Stand-ins: ", "while Dug is down its tools answer with an error, so the agent says so instead of "
                        "improvising from release notes"),
        ("Queue: ", "busy servers line chats up (“you're number 3 in line”)"),
        ("Second example: ", "the same code answers from the FastAPI docs"),
    ]
    s.text(8.6, 2.55, 3.95, 2.7, [[b(k, size=13, color=TWILIGHT, space_before=0 if i == 0 else 9), r(v, size=13)]
                                  for i, (k, v) in enumerate(items)])
    s.notes("""
No screenshot.
- The point for PIs: the assistant isn't BDC-specific code. A new project writes four yaml files (and optionally
  interceptors) and reuses everything else; examples/fastapi proves it with a FastAPI docs bot.
- The BDC config adds the Dug interceptor (Dug rows -> graphs the client draws, plus the studies Dug cites as
  sources) and stand-ins (without them, a down Dug made the agent present release-note study lists as catalog
  results).
- Over capacity, chats wait in line and the client shows their place; the cap was sized with a load test on the
  real backends.
""")


def builder(d):
    s = d.content("bdc-doc-builder: BDC's docs, made searchable",
                  "Runs outside the serving side: it builds the chunks and pushes them, and never touches the database.")
    steps = [
        ("Sources", "sources.yaml lists what to build"),
        ("Load", "files, web pages, git repos, YouTube, transcripts, Freshdesk"),
        ("Chunk", "markdown by heading; videos in timestamped windows"),
        ("Contextualize", "an LLM adds a line of context to each chunk"),
        ("Embed", "bge-m3 vectors, the model doc-mcp queries with"),
        ("Push", "finished records to doc-mcp's ingest API"),
    ]
    sw, gap = 1.75, 0.31
    for i, (name, body) in enumerate(steps):
        x = 0.6 + i * (sw + gap)
        s.box(x + sw / 2 - 0.35, 1.9, 0.7, 0.7, fill=tint(VIOLET, 0.15), line=VIOLET, width=1.75, shape="oval")
        s.text(x + sw / 2 - 0.35, 1.9, 0.7, 0.7, str(i + 1), size=18, color=TWILIGHT, bold=True, align="c",
               anchor="m")
        if i:
            s.line(x - gap - sw / 2 + 0.4, 2.25, x + sw / 2 - 0.4, 2.25, color=VIOLET)
        s.text(x, 2.75, sw, 0.35, name, size=14, color=TWILIGHT, bold=True, align="c")
        s.text(x, 3.12, sw, 1.0, body, size=12, color=INK, align="c")

    cards = [
        ("Incremental", "A re-push embeds only new or changed chunks."),
        ("All or nothing", "A doc type that fails to load is skipped whole; the database keeps its last good build."),
        ("Rehearsable", "A fake, BDC-shaped corpus runs the real CLI against a throwaway database."),
    ]
    for i, (k, v) in enumerate(cards):
        x = 0.6 + i * 4.1
        s.label_box(x, 4.45, 3.85, 1.25, [[b(k, size=14, color=TWILIGHT)], [r(v, size=12, space_before=4)]],
                    fill=SLATE, radius=0.1, pad=0.18)
    s.text(0.6, 6.0, W - 1.2, 0.6,
           [[b("BDC sources: ", size=12, color=TWILIGHT),
             r("the GitBook and website (markdown export), the website FAQ, news, events and fellows, and YouTube "
               "tutorial transcripts.", size=12)]])
    s.notes("""
No screenshot.
- The builder lives outside the serving security boundary: it reads the sources, chunks them, asks an LLM for a
  line of context per chunk (better retrieval), embeds, and pushes finished {id, content, embedding, metadata}
  records to doc-mcp's token-guarded ingest API.
- Pushes are incremental: it looks up what's stored for each source first and only embeds what changed.
- Each doc type builds all or nothing, so a broken source never leaves half a doc type in the database.
- Citations link to the live website and GitBook pages, not the private markdown repo.
""")


def doc_mcp(d):
    s = d.content("bdc-doc-mcp: one search tool over the BDC docs",
                  "The agent calls search_docs; the same search is a plain HTTP API too.")
    # the tool's inputs
    s.box(0.6, 1.95, 5.6, 3.6, fill=tint(ULTRA, 0.12), line=ULTRA, width=1.75, radius=0.1)
    s.text(0.85, 2.1, 5.1, 0.4, "search_docs", size=16, color=TWILIGHT, bold=True)
    params = [
        ("query", "the agent's search words"),
        ("mode", "embedding: by meaning (default)\nkeyword: fuzzy, “picsure” finds PIC-SURE"),
        ("doc_type", "kinds of docs (right); default: the core four"),
        ("dates", "from / to, for news and events"),
    ]
    y = 2.65
    for k, v in params:
        lines = v.count("\n") + 1
        s.text(0.85, y, 1.2, 0.3, k, size=13, color=TWILIGHT, bold=True)
        s.text(2.1, y, 3.95, 0.3 * lines + 0.1, v, size=13)
        y += 0.25 * lines + 0.42

    s.text(6.7, 1.95, 6, 0.35, "Doc types", size=15, color=TWILIGHT, bold=True)
    types = [
        ("docs", "GitBook platform docs", True), ("page", "Key website pages", True),
        ("faq", "Website FAQ", True), ("video", "Tutorial transcripts", True),
        ("update", "News posts (dated)", False), ("event", "Webinars, workshops (dated)", False),
        ("fellow", "BDC Fellows profiles", False),
    ]
    for i, (k, v, default) in enumerate(types):
        x, yy = 6.7 + (i % 2) * 3.05, 2.45 + (i // 2) * 0.82
        s.label_box(x, yy, 2.85, 0.66, [[b(k + "  ", size=13, color=TWILIGHT), r(v, size=11)]], anchor="m",
                    fill=tint(ULTRA, 0.12) if default else WHITE, line=ULTRA, width=1.25, dashed=not default,
                    radius=0.08)
    s.text(9.75, 4.91, 2.85, 0.66, "Solid: searched by default.\nDashed: when asked for.", size=11, color=STONE,
           anchor="m")
    s.text(0.6, 5.85, W - 1.2, 0.6,
           [[b("Also: ", size=12, color=TWILIGHT),
             r("/health reports the document count and each doc type's last push; the vector database sits behind a "
               "five-method interface (Chroma today; pgvector or Qdrant can be swapped in).", size=12)]])
    s.notes("""
No screenshot.
- One MCP tool, search_docs. Embedding mode finds passages by meaning; keyword mode is fuzzy literal matching for
  names and acronyms (picsure, PIC-SURE, Seven Bridges), which embeddings handle poorly.
- Doc types come from a yaml file, so a new deployment declares its own without code changes. The four core
  types are searched by default; news, events and fellows only when the question calls for them.
- Each project gets its own collection, so the BDC, FastAPI and test corpora never mix.
""")


def client_arch(d):
    s = d.content("bdc-assist-client: how an answer becomes UI")
    y1, h1 = 1.85, 1.45
    node(s, 0.6, y1, 2.1, h1, "bdc-assist", "/chat/stream\nevents", "assist")
    node(s, 3.15, y1, 2.85, h1, "bdc-adapter", "Stream events into the\nanswer's parts", "client")
    node(s, 6.45, y1, 2.85, h1, "assistant-ui", "Thread, composer,\nconversation list", "client")
    node(s, 9.75, y1, 3.0, h1, "Answer parts", "Status, text, sources,\nfollow-ups, notes, graph", "client")
    for x1, x2 in [(2.7, 3.15), (6.0, 6.45), (9.3, 9.75)]:
        s.line(x1, y1 + h1 / 2, x2, y1 + h1 / 2)

    node(s, 6.45, 3.85, 2.85, 1.2, "localStorage", "Saved conversations,\nabout 5 MB per site", "client")
    s.line(7.875, y1 + h1, 7.875, 3.85)
    node(s, 9.75, 3.85, 3.0, 1.2, "KnowledgeGraph", "React wrapper: toolbar,\nlegend, details", "client")
    s.line(11.25, y1 + h1, 11.25, 3.85)
    tag(s, 11.35, 3.4, 1.2, "graph part", align="l")

    # the portable library
    s.box(3.15, 5.45, 9.6, 1.35, fill=tint(TWILIGHT, 0.08), line=TWILIGHT, width=1.75, radius=0.1)
    s.text(3.35, 5.55, 2.6, 0.35, "src/kg library", size=13, color=TWILIGHT, bold=True)
    s.text(3.35, 5.92, 2.6, 0.8, "No React, no app imports:\nanother site can embed it", size=11)
    s.label_box(6.1, 5.6, 3.15, 1.05, [[b("model", size=12, color=TWILIGHT)],
                                       [r("merge tool calls, group releases and synonyms, find shared studies",
                                          size=11, space_before=2)]], fill=WHITE, line=TWILIGHT, width=1, pad=0.1)
    s.label_box(9.45, 5.6, 3.15, 1.05, [[b("views", size=12, color=TWILIGHT)],
                                        [r("Force, Radial (Cytoscape, d3-force), Flow (d3-sankey), images",
                                           size=11, space_before=2)]], fill=WHITE, line=TWILIGHT, width=1, pad=0.1)
    s.line(11.25, 5.05, 11.25, 5.45)

    s.text(0.6, 3.85, 5.3, 1.3,
           [[b("Built with ", size=12, color=TWILIGHT),
             r("React 19, TypeScript, Vite, Tailwind and shadcn/ui, assistant-ui for the chat.", size=12)],
            [b("Tested ", size=12, color=TWILIGHT, space_before=8),
             r("against graphs from real Dug results (fixtures generated by bdc-assist).", size=12)]])
    s.notes("""
No screenshot.
- The adapter turns the stream into the answer's parts as they arrive (status lines, draft text, sources, graphs,
  follow-ups, guardrail and service notes); assistant-ui draws the thread and our renderers draw each part.
- Conversations are kept in the browser only, under bdc-assist:* keys in localStorage; a full store opens a dialog
  to delete old ones.
- The graph drawing is a separate library with no framework dependencies (src/kg): pure data transforms plus
  views, themed with CSS variables, so the BDC website or another tool could reuse it.
""")


def life(d):
    s = d.content("Life of a question, as the answer streams")
    stops = [
        ("node", "Workflow step started", "Status line"),
        ("status", "Agent calls a tool", "Status line names the tool"),
        ("token", "Draft answer streams", "Answer text appears (or placeholder bars until checked)"),
        ("sources", "Agent finished", "Source icons and the knowledge graph, while checks still run"),
        ("done", "Final state", "The final answer replaces the draft; follow-ups; guardrail notes"),
        ("saved", "Client", "Conversation saved in the browser"),
    ]
    n = len(stops)
    cw = 1.9
    x0, x1, ly = 0.6 + cw / 2, W - 0.6 - cw / 2, 3.05
    s.line(x0 - 0.5, ly, x1 + 0.5, ly, color=TWILIGHT, width=2)
    step = (x1 - x0) / (n - 1)
    for i, (event, when, sees) in enumerate(stops):
        cx = x0 + i * step
        client = event == "saved"
        s.box(cx - 0.15, ly - 0.15, 0.3, 0.3, fill=WHITE, line=TWILIGHT if client else BLUE, width=2.25,
              shape="oval")
        s.text(cx - cw / 2, 1.75, cw, 0.4, event, size=16, color=TWILIGHT if client else BLUE, bold=True, align="c")
        s.text(cx - cw / 2, 2.15, cw, 0.6, when, size=12, color=STONE, align="c")
        s.label_box(cx - cw / 2, 3.45, cw, 1.45, sees, size=12, align="c", anchor="m",
                    fill=tint(TWILIGHT, 0.08), line=None, radius=0.1)

    s.label_box(0.6, 5.15, 5.9, 1.1,
                [[b("error ", size=13, color=TWILIGHT), r("ends the stream instead of done: the draft is discarded "
                                                           "and the answer says it failed.", size=12)]],
                fill=SLATE, radius=0.1, pad=0.2, anchor="m")
    s.label_box(6.85, 5.15, 5.9, 1.1,
                [[b("blocked ", size=13, color=TWILIGHT), r("in done: a refused question, or a draft replaced by the "
                                                             "reject reply. The client says which.", size=12)]],
                fill=SLATE, radius=0.1, pad=0.2, anchor="m")
    key(s, 0.6, 6.45, [("assist", "server event"), ("client", "client")])
    s.notes("""
No screenshot. This slide bridges into the web client walkthrough.
- Events from /chat/stream in order: node (a workflow step started), status (tool calls), token (draft text;
  reset discards it on a new model turn), sources (sent early with the graphs), done (authoritative final state:
  answer, blocked, follow-ups, unavailable services), or error.
- Sources and graphs arrive before done, so the user can explore them while the output guardrail and follow-up
  steps still run.
- ?reveal=after-check holds the draft behind placeholder bars until the output guardrail has passed it, for
  deployments that don't want unchecked text on screen.
""")


def build():
    d = Deck()
    title_slide(d)
    s = d.divider(1, "The system", "Three services and a web client: how a question gets answered.")
    s.notes("Section divider. No screenshot.")
    architecture(d)
    workflow(d)
    template(d)
    builder(d)
    doc_mcp(d)
    client_arch(d)
    life(d)
    d.save()
    print(f"wrote {OUT.relative_to(HERE.parent)} ({len(d.slides)} slides)")
    if "--preview" in sys.argv:
        preview(d)


# ---- preview ------------------------------------------------------------------------------------------------


def preview(d, dpi=100):
    """Rough PNGs of each slide, to check layout. Warns where text overflows its box."""
    from PIL import Image, ImageDraw, ImageFont

    fonts = {}

    def font(size, bold):
        k = (size, bold)
        if k not in fonts:
            name = "LiberationSans-Bold.ttf" if bold else "LiberationSans-Regular.ttf"
            fonts[k] = ImageFont.truetype(f"/usr/share/fonts/truetype/liberation/{name}", round(size * dpi / 72))
        return fonts[k]

    def px(v):
        return round(v * dpi)

    def hexc(h):
        return "#" + h

    PREVIEW.mkdir(exist_ok=True)
    for old in PREVIEW.glob("slide-*.png"):
        old.unlink()
    for n, sl in enumerate(d.slides, 1):
        img = Image.new("RGB", (px(W), px(H)), "white")
        g = ImageDraw.Draw(img)
        for op in sl.ops:
            kind = op[0]
            if kind == "bg":
                g.rectangle([0, 0, px(W), px(H)], fill=hexc(op[1]))
            elif kind == "box":
                _, x, y, w, h, fill, line, width, dashed, radius, shape = op
                xy = [px(x), px(y), px(x + w), px(y + h)]
                kw = {"fill": hexc(fill) if fill else None, "outline": hexc(line) if line else None,
                      "width": max(1, round(width * dpi / 72))}
                if shape == "oval":
                    g.ellipse(xy, **kw)
                elif shape == "rect":
                    g.rectangle(xy, **kw)
                else:
                    g.rounded_rectangle(xy, radius=px(radius), **kw)
                if dashed and line:  # dashes: knock gaps out of the outline
                    gap = (fill and hexc(fill)) or hexc(sl.bg)
                    for t in range(xy[0] + 6, xy[2] - 6, 12):
                        g.line([t, xy[1], t + 4, xy[1]], fill=gap, width=kw["width"] + 1)
                        g.line([t, xy[3], t + 4, xy[3]], fill=gap, width=kw["width"] + 1)
                    for t in range(xy[1] + 6, xy[3] - 6, 12):
                        g.line([xy[0], t, xy[0], t + 4], fill=gap, width=kw["width"] + 1)
                        g.line([xy[2], t, xy[2], t + 4], fill=gap, width=kw["width"] + 1)
            elif kind == "line":
                _, x1, y1, x2, y2, color, width, arrow, dashed = op
                wpx = max(1, round(width * dpi / 72))
                g.line([px(x1), px(y1), px(x2), px(y2)], fill=hexc(color), width=wpx)
                if arrow:
                    import math

                    a = math.atan2(y2 - y1, x2 - x1)
                    L = 9
                    p = [(px(x2), px(y2)),
                         (px(x2) - L * math.cos(a - 0.45), px(y2) - L * math.sin(a - 0.45)),
                         (px(x2) - L * math.cos(a + 0.45), px(y2) - L * math.sin(a + 0.45))]
                    g.polygon(p, fill=hexc(color))
            elif kind == "image":
                _, file, x, y, w, h = op
                img.paste(Image.open(file).convert("RGB").resize((px(w), px(h))), (px(x), px(y)))
            elif kind == "text":
                _, x, y, w, h, paras, align, anchor, pad, spacing = op
                inner = px(w - 2 * pad)
                lines = []  # (runs, height)
                for para in paras:
                    words = [(word, o) for t, o in para for word in _split(t)]
                    line, lw, lh = [], 0, 0
                    before = para[0][1].get("space_before", 0) * dpi / 72 if para else 0
                    first = True
                    for word, o in words:
                        f = font(o["size"], o["bold"])
                        ww = g.textlength(word, font=f)
                        if g.textlength(word.strip(), font=f) > inner + 1:
                            print(f"  slide {n}: word wider than its box: {word.strip()!r}")
                        if line and lw + ww > inner and word.strip():
                            lines.append((line, lh, before if first else 0))
                            first = False
                            line, lw = [], 0
                            if not word.strip():
                                continue
                        line.append((word, o, f))
                        lw += ww
                        lh = max(lh, o["size"] * 1.2 * spacing * dpi / 72)
                    if not line:
                        lh = para[0][1]["size"] * 1.2 * dpi / 72 if para else 14
                    lines.append((line, lh, before if first else 0))
                total = sum(lh + bf for _, lh, bf in lines)
                avail = px(h - 2 * pad)
                if total > avail + 2:
                    text = "".join(t for para in paras for t, _ in para)[:50]
                    print(f"  slide {n}: text overflows by {(total - avail) / dpi:.2f} in: {text!r}")
                ty = px(y + pad) + {"t": 0, "m": (avail - total) / 2, "b": avail - total}[anchor]
                for line, lh, bf in lines:
                    ty += bf
                    lw = sum(g.textlength(wd, font=f) for wd, _, f in line)
                    tx = px(x + pad) + {"l": 0, "c": (inner - lw) / 2, "r": inner - lw}[align]
                    for wd, o, f in line:
                        g.text((tx, ty + lh - o["size"] * 1.2 * dpi / 72), wd, font=f, fill=hexc(o["color"]))
                        tx += g.textlength(wd, font=f)
                    ty += lh
        out = PREVIEW / f"slide-{n:02d}.png"
        img.save(out)
    print(f"preview in {PREVIEW.relative_to(HERE.parent)}/")


def _split(t):
    """Words with the spaces after them, so lines wrap between words."""
    import re

    return re.findall(r"\S+\s*|\s+", t)


if __name__ == "__main__":
    build()
